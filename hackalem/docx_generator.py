"""Түзетілген деректерден DOCX хаттамасын құрастыру."""

import io

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from utils import clean_text, normalize_tasks


def generate_protocol_docx(
    meeting_title: str, meeting_date: str, summary: str,
    speakers: list, tasks: list, transcript_text: str | None = None,
) -> io.BytesIO:
    rows = normalize_tasks(tasks)
    if speakers is None:
        speakers = []
    if not isinstance(speakers, (list, tuple)):
        raise ValueError("Сөйлеушілер тізім түрінде берілуі керек.")
    names = list(dict.fromkeys(name for item in speakers if (name := clean_text(item, ""))))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name in ("Normal", "Title", "Heading 1"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(18 if name == "Title" else 13 if name == "Heading 1" else 12)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)

    paragraph = doc.add_paragraph(clean_text(meeting_title, "Кеңес хаттамасы"), "Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph = doc.add_paragraph(f"Күні: {clean_text(meeting_date)}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("Мәтінде расталған сөйлеушілер", level=1)
    for name in names or ["Анықталмады"]:
        doc.add_paragraph(name)
    doc.add_heading("Қысқаша қорытынды", level=1)
    doc.add_paragraph(clean_text(summary))
    doc.add_heading("Тапсырмалар", level=1)

    if rows:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = [Cm(0.8), Cm(8.2), Cm(4.5), Cm(3.5)]
        for column, width in zip(table.columns, widths):
            column.width = width
        for cell, text in zip(table.rows[0].cells, ["№", "Тапсырма", "Жауапты тұлға", "Мерзім"]):
            cell.text = text
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "D9E2F3")
            cell._tc.get_or_add_tcPr().append(shade)
            for run in cell.paragraphs[0].runs:
                run.bold = True
            cell.paragraphs[0].paragraph_format.keep_with_next = True
        header = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(header)
        for number, task in enumerate(rows, 1):
            cells = table.add_row().cells
            values = [str(number), task["task"], task["responsible"], task["deadline"]]
            for cell, value in zip(cells, values):
                cell.text = value
        for row in table.rows:
            for index, (cell, width) in enumerate(zip(row.cells, widths)):
                cell.width = width
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(4)
                    paragraph.paragraph_format.space_before = Pt(4)
                    if index == 0:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        doc.add_paragraph("Тапсырмалар табылмады.")

    transcript = clean_text(transcript_text, "")
    if transcript:
        doc.add_page_break()
        doc.add_heading("Толық транскрипт", level=1)
        for line in transcript.splitlines():
            paragraph = doc.add_paragraph(line)
            for run in paragraph.runs:
                run.font.size = Pt(10)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
