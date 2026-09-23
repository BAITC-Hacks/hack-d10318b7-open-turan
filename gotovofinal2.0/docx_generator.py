"""Создание редактируемого DOCX-протокола после проверки человеком."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from utils import UNKNOWN, clean_text, normalize_tasks


def _set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def _keep_row_together(row) -> None:
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def _set_repeat_table_header(row) -> None:
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(header)


def _format_cell(cell, font_size: int = 10) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.space_before = Pt(3)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(font_size)


def generate_protocol_docx(
    *,
    meeting_title: str,
    meeting_date: str,
    summary: str,
    decisions: list[str] | None,
    speakers: list[str] | None,
    tasks: list[dict],
    transcript_text: str | None = None,
    output_language: str = "ru",
) -> BytesIO:
    kk = output_language == "kk"
    def label(ru, kz):
        return kz if kk else ru
    rows = normalize_tasks(tasks)
    title_text = clean_text(meeting_title, "Meeting protocol")
    date_text = clean_text(meeting_date, UNKNOWN)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)

    for name, size in (("Normal", 10), ("Title", 20), ("Heading 1", 14), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)

    title = doc.add_paragraph(title_text, style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date = doc.add_paragraph(f"{label('Дата встречи', 'Кездесу күні')}: {date_text}")
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading(label("Краткое резюме", "Қысқаша қорытынды"), level=1)
    doc.add_paragraph(clean_text(summary))

    doc.add_heading(label("Участники, явно упомянутые в речи", "Сөз барысында аталған қатысушылар"), level=1)
    names = [clean_text(name, "") for name in (speakers or [])]
    names = list(dict.fromkeys(name for name in names if name))
    doc.add_paragraph(", ".join(names) if names else UNKNOWN)

    doc.add_heading(label("Принятые решения", "Қабылданған шешімдер"), level=1)
    clean_decisions = [clean_text(item, "") for item in (decisions or [])]
    clean_decisions = [item for item in clean_decisions if item]
    if clean_decisions:
        for item in clean_decisions:
            doc.add_paragraph(item, style="List Bullet")
    else:
        doc.add_paragraph(label("Явные решения не зафиксированы.", "Нақты шешімдер тіркелмеген."))

    doc.add_heading(label("Задачи", "Тапсырмалар"), level=1)
    if rows:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = [Cm(0.8), Cm(7.7), Cm(4.0), Cm(3.2)]
        for column, width in zip(table.columns, widths):
            column.width = width
        headers = ["№", label("Задача и подтверждение", "Тапсырма және дәйексөз"),
                   label("Ответственный", "Жауапты"), label("Срок", "Мерзімі")]
        for cell, header_text in zip(table.rows[0].cells, headers):
            cell.text = header_text
            _set_cell_shading(cell, "D9E2F3")
            _format_cell(cell, 10)
            for run in cell.paragraphs[0].runs:
                run.bold = True
            if header_text == "№":
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_repeat_table_header(table.rows[0])
        _keep_row_together(table.rows[0])

        for number, task in enumerate(rows, start=1):
            cells = table.add_row().cells
            values = [str(number), "", task["responsible"], task["deadline"]]
            for cell, width, value in zip(cells, widths, values):
                cell.width = width
                cell.text = value
                _format_cell(cell)
            cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            task_paragraph = cells[1].paragraphs[0]
            task_run = task_paragraph.add_run(task["task"])
            task_run.bold = True
            evidence = task.get("evidence", "")
            if evidence:
                citation = cells[1].add_paragraph()
                citation.paragraph_format.space_before = Pt(2)
                citation_run = citation.add_run(f"{label('Подтверждение', 'Дәйексөз')}: {evidence}")
                citation_run.italic = True
                citation_run.font.size = Pt(9)
            _keep_row_together(table.rows[-1])
    else:
        doc.add_paragraph(label("Задач не найдено. При необходимости добавьте их вручную.",
                               "Тапсырмалар табылмады. Қажет болса, қолмен қосыңыз."))

    transcript = clean_text(transcript_text, "")
    if transcript:
        doc.add_page_break()
        doc.add_heading(label("Полная расшифровка", "Толық мәтін"), level=1)
        for line in transcript.splitlines():
            if line.strip():
                paragraph = doc.add_paragraph(line.strip())
                paragraph.paragraph_format.keep_together = False
                for run in paragraph.runs:
                    run.font.size = Pt(9)

    doc.core_properties.title = title_text
    doc.core_properties.subject = "Проект протокола совещания"
    doc.core_properties.author = "Meeting Notes"
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output
