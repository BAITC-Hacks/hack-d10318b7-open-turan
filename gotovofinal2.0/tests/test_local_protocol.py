import unittest
from unittest.mock import patch
from pathlib import Path
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from pydantic import ValidationError

from docx_generator import generate_protocol_docx
from main import app
import main
from utils import (
    MeetingAnalysis,
    MeetingTask,
    ProcessingError,
    clean_text,
    normalize_tasks,
    validate_local_ollama_url,
    Transcript,
    ollama_status,
)


class TextValidationTests(unittest.TestCase):
    def test_clean_text_removes_xml_controls_and_empty_values(self):
        self.assertEqual(clean_text("Имя\x00 участника"), "Имя участника")
        self.assertEqual(clean_text("NaN"), "Не определено")
        self.assertEqual(clean_text(None, ""), "")

    def test_normalize_tasks_drops_empty_rows_and_fills_missing_fields(self):
        tasks = normalize_tasks(
            [
                {"task": "   ", "responsible": "", "deadline": "", "evidence": ""},
                {"task": "Отправить итог", "responsible": None, "deadline": float("nan")},
            ]
        )
        self.assertEqual(len(tasks), 1)
        self.assertEqual(tasks[0]["task"], "Отправить итог")
        self.assertEqual(tasks[0]["responsible"], "Не определено")
        self.assertEqual(tasks[0]["deadline"], "Не определено")
        self.assertEqual(tasks[0]["evidence"], "")

    def test_meeting_schema_requires_an_evidence_quote(self):
        with self.assertRaises(ValidationError):
            MeetingAnalysis.model_validate(
                {"status": "ok", "summary": "Кратко", "decisions": [], "speakers": [], "tasks": [
                    {"task": "Сделать", "responsible": "Не определено", "deadline": "Не определено"}
                ]}
            )

    def test_remote_ollama_url_is_rejected(self):
        with self.assertRaises(ProcessingError):
            validate_local_ollama_url("https://example.com")


class DocumentTests(unittest.TestCase):
    def test_docx_contains_reviewed_task_and_transcript(self):
        output = generate_protocol_docx(
            meeting_title="Встреча команды",
            meeting_date="2026-09-23",
            summary="Обсудили запуск.",
            decisions=["Запускать после проверки."],
            speakers=["Айжан"],
            tasks=[
                MeetingTask(
                    task="Подготовить план",
                    responsible="Айжан",
                    deadline="Пятница",
                    evidence="Айжан подготовит план к пятнице",
                )
            ],
            transcript_text="Айжан подготовит план к пятнице.",
        )
        self.assertIsInstance(output, BytesIO)
        document = Document(output)
        paragraphs = "\n".join(item.text for item in document.paragraphs)
        table_text = "\n".join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        )
        self.assertIn("Встреча команды", paragraphs)
        self.assertIn("Подготовить план", table_text)
        self.assertIn("подготовит план", table_text)
        self.assertNotIn("None", paragraphs + table_text)


class ApiSmokeTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def test_home_serves_the_combined_recorder_and_review_ui(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Начать запись", response.text)
        self.assertIn("Добавить задачу вручную", response.text)

    def test_health_reports_local_processing_without_cloud_credentials(self):
        response = self.client.get("/health")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["privacy"], "local_only")
        self.assertIn("ollama_online", response.json())

    def test_rejects_unsupported_upload_and_unknown_document(self):
        rejected = self.client.post(
            "/transcribe",
            files={"file": ("note.exe", b"bad", "application/octet-stream")},
            data={"meeting_title": "Тест", "meeting_date": "2026-09-23"},
        )
        self.assertEqual(rejected.status_code, 400)
        missing = self.client.post(
            "/transcribe/not-a-task/document",
            json={
                "meeting_title": "Тест",
                "meeting_date": "2026-09-23",
                "summary": "Проверка маршрута.",
                "decisions": [],
                "speakers": [],
                "tasks": [],
            },
        )
        self.assertEqual(missing.status_code, 404)


class RegressionTests(unittest.TestCase):
    def setUp(self):
        main.tasks.clear()
        self.client = TestClient(app)
        self.metadata = {"meeting_title": "Команда", "meeting_date": "2026-09-23", "output_language": "kk"}

    def test_text_survives_unavailable_model_and_exports_reviewed_kazakh_docx(self):
        with patch("main.structure_meeting", side_effect=ProcessingError("Ollama недоступна")):
            submitted = self.client.post("/transcript", json={**self.metadata, "transcript_text": "Айжан подготовит план к пятнице."})
        self.assertEqual(submitted.status_code, 202)
        task_id = submitted.json()["task_id"]
        result = self.client.get(f"/transcribe/{task_id}").json()["result"]
        self.assertEqual(result["status"], "manual")
        self.assertEqual(result["summary"], "")
        self.assertIn("Айжан", result["transcript_text"])
        exported = self.client.post(f"/transcribe/{task_id}/document", json={
            **self.metadata, "summary": "Жоспар талқыланды.", "tasks": [
                {"task": "Жоспар дайындау", "responsible": "Айжан", "evidence": "Айжан подготовит план к пятнице"}],
            "transcript_text": result["transcript_text"]})
        self.assertEqual(exported.status_code, 200)
        doc = Document(BytesIO(exported.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Қысқаша қорытынды", text)
        self.assertIn("Толық мәтін", text)
        self.assertIn("Жоспар дайындау", doc.tables[0].cell(1, 1).text)

    def test_required_fields_reject_whitespace_and_xml_controls(self):
        for field in ("meeting_title", "summary"):
            body = {**self.metadata, "summary": "Итоги", field: " \x00 "}
            self.assertEqual(self.client.post("/transcribe/missing/document", json=body).status_code, 422)
        self.assertEqual(self.client.post("/transcript", json={**self.metadata, "transcript_text": "  "}).status_code, 422)

    def test_rejects_invalid_date_and_task(self):
        response = self.client.post("/transcript", json={**self.metadata, "meeting_date": "2026-02-31", "transcript_text": "Текст"})
        self.assertEqual(response.status_code, 422)
        response = self.client.post("/transcribe/missing/document", json={**self.metadata, "summary": "Итоги", "tasks": [{"task": "  "}]})
        self.assertEqual(response.status_code, 422)

    def test_empty_and_oversized_uploads_do_not_reserve_tasks(self):
        for content in (b"", b"12345"):
            with patch("main.MAX_FILE_SIZE_BYTES", 4):
                response = self.client.post("/transcribe", files={"file": ("x.wav", content)}, data=self.metadata)
            self.assertEqual(response.status_code, 400)
        self.assertEqual(main.tasks, {})

    def test_audio_fallback_cleans_temporary_file(self):
        captured = []
        def transcribe(path, language):
            self.assertTrue(Path(path).exists())
            captured.append(path)
            return Transcript("Айжан подготовит план.", [], "")
        with patch("main.transcribe_audio", side_effect=transcribe), patch("main.structure_meeting", side_effect=ProcessingError("Нет модели")):
            response = self.client.post("/transcribe", files={"file": ("audio.wav", b"test")}, data=self.metadata)
        self.assertEqual(response.status_code, 202)
        self.assertFalse(Path(captured[0]).exists())
        result = self.client.get("/transcribe/" + response.json()["task_id"]).json()
        self.assertEqual(result["status"], "done")
        self.assertIn("Айжан", result["result"]["transcript_text"])

    def test_active_limit_and_completed_eviction_preserve_active_tasks(self):
        first = main.reserve_task()
        second = main.reserve_task()
        with self.assertRaises(main.HTTPException):
            main.reserve_task()
        main.update_task(first, status="done", result={})
        with patch("main.MAX_TASKS_IN_MEMORY", 2):
            third = main.reserve_task()
        self.assertEqual(set(main.tasks), {second, third})

    def test_health_requires_both_components(self):
        with patch("main.ollama_status", return_value={"online": True, "model_available": True}), patch("main.importlib.util.find_spec", return_value=None):
            response = self.client.get("/health").json()
        self.assertEqual(response["status"], "setup_required")
        self.assertFalse(response["transcription_ready"])

    def test_malformed_ollama_health_is_handled(self):
        for payload in ([], {"models": None}, {"models": [{"name": []}]}):
            with patch("utils.httpx.get") as get:
                get.return_value.json.return_value = payload
                self.assertFalse(ollama_status()["model_available"])

    def test_empty_task_with_responsible_is_not_exported(self):
        self.assertEqual(normalize_tasks([{"task": " ", "responsible": "Айжан"}]), [])

    def test_static_assets_are_served_locally(self):
        for filename in ("app.js", "style.css"):
            self.assertEqual(self.client.get("/static/" + filename).status_code, 200)


if __name__ == "__main__":
    unittest.main()
