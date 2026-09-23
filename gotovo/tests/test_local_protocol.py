import unittest
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from pydantic import ValidationError

from docx_generator import generate_protocol_docx
from main import app
from utils import (
    MeetingAnalysis,
    MeetingTask,
    ProcessingError,
    clean_text,
    normalize_tasks,
    validate_local_ollama_url,
)


class TextValidationTests(unittest.TestCase):
    def test_clean_text_removes_xml_controls_and_empty_values(self):
        self.assertEqual(clean_text("Имя\x00 участника"), "Имя участника")
        self.assertEqual(clean_text("NaN"), "Не определено")
        self.assertEqual(clean_text(None, ""), "")

    def test_normalize_tasks_drops_empty_rows_and_fills_missing_fields(self):
        tasks = normalize_tasks(
            [
                {"task": "   ", "responsible": "", "deadline": "", "evidence": ""},
                {"task": "Отправить итог", "responsible": None, "deadline": float("nan")},
            ]
        )
        self.assertEqual(len(tasks), 1)
        self.assertEqual(tasks[0]["task"], "Отправить итог")
        self.assertEqual(tasks[0]["responsible"], "Не определено")
        self.assertEqual(tasks[0]["deadline"], "Не определено")
        self.assertEqual(tasks[0]["evidence"], "")

    def test_meeting_schema_requires_an_evidence_quote(self):
        with self.assertRaises(ValidationError):
            MeetingAnalysis.model_validate(
                {"status": "ok", "summary": "Кратко", "decisions": [], "speakers": [], "tasks": [
                    {"task": "Сделать", "responsible": "Не определено", "deadline": "Не определено"}
                ]}
            )

    def test_remote_ollama_url_is_rejected(self):
        with self.assertRaises(ProcessingError):
            validate_local_ollama_url("https://example.com")


class DocumentTests(unittest.TestCase):
    def test_docx_contains_reviewed_task_and_transcript(self):
        output = generate_protocol_docx(
            meeting_title="Встреча команды",
            meeting_date="2026-09-23",
            summary="Обсудили запуск.",
            decisions=["Запускать после проверки."],
            speakers=["Айжан"],
            tasks=[
                MeetingTask(
                    task="Подготовить план",
                    responsible="Айжан",
                    deadline="Пятница",
                    evidence="Айжан подготовит план к пятнице",
                )
            ],
            transcript_text="Айжан подготовит план к пятнице.",
        )
        self.assertIsInstance(output, BytesIO)
        document = Document(output)
        paragraphs = "\n".join(item.text for item in document.paragraphs)
        table_text = "\n".join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        )
        self.assertIn("Встреча команды", paragraphs)
        self.assertIn("Подготовить план", table_text)
        self.assertIn("подготовит план", table_text)
        self.assertNotIn("None", paragraphs + table_text)


class ApiSmokeTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def test_home_serves_the_combined_recorder_and_review_ui(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Начать запись", response.text)
        self.assertIn("Добавить задачу вручную", response.text)

    def test_health_reports_local_processing_without_cloud_credentials(self):
        response = self.client.get("/health")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["privacy"], "local_only")
        self.assertIn("ollama_online", response.json())

    def test_rejects_unsupported_upload_and_unknown_document(self):
        rejected = self.client.post(
            "/transcribe",
            files={"file": ("note.exe", b"bad", "application/octet-stream")},
            data={"meeting_title": "Тест", "meeting_date": "2026-09-23"},
        )
        self.assertEqual(rejected.status_code, 400)
        missing = self.client.post(
            "/transcribe/not-a-task/document",
            json={
                "meeting_title": "Тест",
                "meeting_date": "2026-09-23",
                "summary": "Проверка маршрута.",
                "decisions": [],
                "speakers": [],
                "tasks": [],
            },
        )
        self.assertEqual(missing.status_code, 404)


if __name__ == "__main__":
    unittest.main()

